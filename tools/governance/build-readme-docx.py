from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "README.md"
TARGET = ROOT / "README.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def keep_row_together(row):
    properties = row._tr.get_or_add_trPr()
    cant_split = properties.find(qn("w:cantSplit"))
    if cant_split is None:
        properties.append(OxmlElement("w:cantSplit"))


def mark_header_row(row):
    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        properties.append(OxmlElement("w:tblHeader"))


def set_table_geometry(table, width_twips):
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_twips))

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    table.autofit = False


def plain(text):
    text = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "").replace("**", "")
    return text.strip()


def add_text(doc, text, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(plain(text))
    return paragraph


def main():
    source_text = SOURCE.read_text(encoding="utf-8")
    lines = source_text.splitlines()
    version = (ROOT / "VERSION").read_text(encoding="utf-8").strip()
    batch_match = re.search(r"complete through batch `([^`]+)`", source_text)
    closed_batch = batch_match.group(1) if batch_match else "current"

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    content_width_twips = round(
        (
            section.page_width
            - section.left_margin
            - section.right_margin
        )
        / 635
    )

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.25)
    styles["Normal"].paragraph_format.space_after = Pt(3.5)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 1"].font.color.rgb = RGBColor(49, 98, 115)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Colonist Awareness Overhaul")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(49, 98, 115)
    subtitle = doc.add_paragraph("RimWorld 1.6 framework and simulation overhaul")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.color.rgb = RGBColor(89, 89, 89)

    i = 1
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            add_text(doc, line[3:], "Heading 1")
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            body = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            run = paragraph.add_run("\n".join(body))
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(8.5)
            if language:
                run.font.color.rgb = RGBColor(38, 88, 105)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [plain(value) for value in lines[i].strip().strip("|").split("|")]
                rows.append(cells)
                i += 1
            rows = [row for row in rows if not all(re.fullmatch(r"-+", c or "") for c in row)]
            table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            set_table_geometry(table, content_width_twips)
            for r_index, row in enumerate(rows):
                keep_row_together(table.rows[r_index])
                if r_index == 0:
                    mark_header_row(table.rows[r_index])
                for c_index, value in enumerate(row):
                    cell = table.cell(r_index, c_index)
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Aptos"
                            run.font.size = Pt(8)
                            if r_index == 0:
                                run.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                    if r_index == 0:
                        set_cell_shading(cell, "316273")
                    elif r_index % 2 == 0:
                        set_cell_shading(cell, "EAF1F3")
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                add_text(doc, lines[i].strip()[2:], "List Bullet")
                i += 1
            continue

        paragraph_lines = [line]
        i += 1
        while i < len(lines):
            candidate = lines[i].strip()
            if (not candidate or candidate.startswith("#") or candidate.startswith("|")
                    or candidate.startswith("-") or candidate.startswith("```")
                    or candidate.startswith("<!--")):
                break
            paragraph_lines.append(candidate)
            i += 1
        add_text(doc, " ".join(paragraph_lines))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(
        f"Colonist Awareness Overhaul | {version} | {closed_batch}"
    ).font.size = Pt(8)
    doc.core_properties.title = "Colonist Awareness Overhaul"
    doc.core_properties.subject = "Current project README"
    doc.core_properties.author = "ellyj3rain"
    doc.core_properties.keywords = (
        f"RimWorld, Colonist Awareness, {closed_batch}, Culture"
    )
    doc.save(TARGET)


if __name__ == "__main__":
    main()
